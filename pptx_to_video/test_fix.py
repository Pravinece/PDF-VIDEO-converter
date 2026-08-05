"""Test pronunciation fixes on the actual PPTX voiceover notes."""
import re
import sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation

PRONUNCIATION_MAP = {
    "Pradhan Mantri": "Pruh-dhaan Mun-tree",
    "Mudra": "Mood-raa",
    "Yojna": "Yoj-naa",
    "Yojana": "Yoj-naa",
    "PMMY": "P M M Y",
    "Shishu": "Shee-shoo",
    "Kishor": "Kish-ore",
    "Kishore": "Kish-ore",
    "Tarun": "Tuh-roon",
    "PNB": "P N B",
    "MSME": "M S M E",
    "CGTMSE": "C G T M S E",
    "Udyog Aadhaar": "Ud-yog Aa-dhaar",
    "DIGI MSME": "Digi M S M E",
    "lakh": "laakh",
}


def fix(text):
    text = re.sub(r'^(voiceover|narration|voiceover\s*/\s*narration)\s*:\s*', '', text, flags=re.IGNORECASE)
    text = re.sub(r'(?<![\d])/(?![\d])', ' or ', text)
    for term, phonetic in PRONUNCIATION_MAP.items():
        text = re.compile(re.escape(term), re.IGNORECASE).sub(phonetic, text)
    return text.strip()


pptx_path = r"C:\Users\s4028\Downloads\Revised Modules\Module No. 1 - Govt Sponsored Credit Schemes\Topic 1- Pradhan Mantri Mudra Yojana\PNB-Pradhan Mantri Mudra Yojana.pptx"
prs = Presentation(pptx_path)

print("=" * 80)
print("AFTER PRONUNCIATION FIX - What TTS will receive:")
print("=" * 80)

for i, slide in enumerate(prs.slides):
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            fixed = fix(notes)
            print(f"\n--- Slide {i+1} ---")
            print(fixed[:300])

print("\n" + "=" * 80)
